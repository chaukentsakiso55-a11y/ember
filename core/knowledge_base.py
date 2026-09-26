"""Offline knowledge library for EMBER with lightweight local indexing."""
from __future__ import annotations
import json, re, shutil
from pathlib import Path
from datetime import datetime

BASE=Path(__file__).resolve().parent.parent; ROOT=BASE/'knowledge'; INDEX=ROOT/'index.json'
def _load():
    try:return json.loads(INDEX.read_text(encoding='utf-8'))
    except Exception:return {'documents':{}}
def _save(d):ROOT.mkdir(parents=True,exist_ok=True);INDEX.write_text(json.dumps(d,indent=2,ensure_ascii=False),encoding='utf-8')
def _text(path:Path)->str:
    ext=path.suffix.lower()
    if ext in {'.txt','.md','.py','.json','.csv','.log','.html','.htm','.js','.kt','.java','.xml','.yaml','.yml'}:
        return path.read_text(encoding='utf-8',errors='replace')
    if ext=='.pdf':
        try:
            try:from pypdf import PdfReader
            except Exception:from PyPDF2 import PdfReader
            return '\n'.join((p.extract_text() or '') for p in PdfReader(str(path)).pages)
        except Exception as e:return f'[PDF extraction unavailable: {e}]'
    if ext=='.docx':
        try:
            from docx import Document
            return '\n'.join(p.text for p in Document(str(path)).paragraphs)
        except Exception as e:return f'[DOCX extraction unavailable: {e}]'
    if ext=='.pptx':
        try:
            from pptx import Presentation
            return '\n'.join(shape.text for slide in Presentation(str(path)).slides for shape in slide.shapes if hasattr(shape,'text'))
        except Exception as e:return f'[PPTX extraction unavailable: {e}]'
    if ext in {'.xlsx','.xlsm'}:
        try:
            from openpyxl import load_workbook
            wb=load_workbook(str(path),read_only=True,data_only=True);rows=[]
            for ws in wb.worksheets:
                rows.append(f'[{ws.title}]')
                for row in ws.iter_rows(values_only=True):rows.append(' | '.join('' if v is None else str(v) for v in row))
            return '\n'.join(rows)
        except Exception as e:return f'[Spreadsheet extraction unavailable: {e}]'
    return ''
def add(path:str,title:str='')->str:
    src=Path(path).expanduser().resolve()
    if not src.is_file():return 'File not found.'
    ROOT.mkdir(parents=True,exist_ok=True);dst=ROOT/src.name
    if src!=dst:shutil.copy2(src,dst)
    text=_text(dst)
    if not text.strip() or text.startswith('[') and 'unavailable:' in text:return text or 'That file type could not be indexed.'
    data=_load();data.setdefault('documents',{})[dst.name]={'title':title or src.stem,'path':str(dst),'text':text[:2_000_000],'added':datetime.now().isoformat(timespec='seconds')};_save(data)
    return f"Indexed '{title or src.stem}' ({len(text):,} characters)."
def remove(name:str)->bool:
    d=_load();docs=d.get('documents',{}); key=next((k for k,v in docs.items() if k==name or str(v.get('title','')).casefold()==str(name).casefold()),None)
    if not key:return False
    rec=docs.pop(key);_save(d)
    try:Path(rec.get('path','')).unlink(missing_ok=True)
    except Exception:pass
    return True
def list_docs():return [{k:v.get('title',k)} for k,v in _load().get('documents',{}).items()]
def search(query:str,limit:int=8)->str:
    original=str(query).casefold()
    # Normalize recurring speech-to-text/project-name variants before tokenization.
    phrase_aliases={
        'study log':'studylock','studylog':'studylock','starlock':'studylock','study lock':'studylock',
        'cyber powers':'cyber pulse','cyberpowers':'cyber pulse','cyberpulse':'cyber pulse',
        'study os':'studyos','study ai':'studyai','star ai':'starai','guardian ai':'guardianai',
        'pulsar ai':'pulsarai','pulsa ai':'pulsarai','pulsar core':'pulsarai',
        'infinity prime':'infinityprime','infinity os':'infinityos','alpha eleven':'alpha11'
    }
    normalized=original
    for a,b in phrase_aliases.items():normalized=normalized.replace(a,b)
    raw=re.findall(r'[A-Za-z0-9]{2,}',normalized)
    stop={'a','an','and','are','as','at','be','been','by','can','did','do','does','for','from','how','i','in','is','it','me','of','on','or','tell','that','the','their','them','there','these','they','this','to','was','were','what','when','where','which','who','why','with','you','behind','about','explain','information','info','part','product','belong','belongs','member','owned','owns'}
    q=[t for t in raw if t not in stop]
    if not q:q=raw
    if not q:return 'Provide a search query.'
    hits=[]
    normalized_phrase=' '.join(q)
    compact_phrase=''.join(q)
    for fn,doc in _load().get('documents',{}).items():
        text=doc.get('text','');low=text.casefold();title=str(doc.get('title',fn));title_low=title.casefold();score=0.0
        title_norm=title_low
        for a,b in phrase_aliases.items():title_norm=title_norm.replace(a,b)
        title_tokens=set(re.findall(r'[a-z0-9]{2,}',title_norm))
        compact_title=title_norm.replace(' ','').replace('_','').replace('-','')
        matched=[]
        for t in q:
            c=low.count(t)
            if c:
                matched.append(t);score += min(c,12) * min(6.0, 1.0 + len(t)/2.0)
            if t in title_tokens or t in title_norm or t in compact_title:score += 24.0
        if q and all(t in title_norm for t in q):score += 70.0
        if normalized_phrase and normalized_phrase in low:score += 35.0
        if compact_phrase and compact_phrase in compact_title:score += 90.0
        if title.startswith('PROJECT_OWNERSHIP_STATUS') and any(t in compact_title for t in q):score += 150.0
        # Exact project-name section files should outrank broad master documents.
        if title.startswith('MASTER_') and q and all(t in title_norm for t in q):score += 120.0
        if score:
            anchor=max(matched,key=len) if matched else q[0]
            pos=low.find(anchor)
            if pos < 0:pos=0
            start=max(0,pos-280);snippet=text[start:start+1100].replace('\n',' ')
            hits.append((score,len(matched),title,snippet))
    hits.sort(key=lambda x:(x[0],x[1]),reverse=True)
    return '\n\n'.join(f'[{title}] {snippet}' for _,__,title,snippet in hits[:limit]) or 'No local knowledge matched that query.'
